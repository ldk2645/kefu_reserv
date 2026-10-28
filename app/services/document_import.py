import hashlib
from pathlib import Path
import re
from uuid import uuid4
from datetime import date

from app.repositories.catalog import Catalog, now

KINDS = {"official_manual": "官方说明书", "official_faq": "官方 FAQ",
         "internal_faq": "内部已审核 FAQ", "shop_policy": "店铺 SOP / 平台规则"}


def extract_pages(path: Path):
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md"}:
            return [(None, path.read_text(encoding="utf-8-sig"))]
        if suffix == ".docx":
            from docx import Document
            doc = Document(path)
            text = "\n\n".join(p.text for p in doc.paragraphs)
            text += "\n\n" + "\n".join(" | ".join(c.text for c in r.cells)
                                          for t in doc.tables for r in t.rows)
            return [(None, text)]
        if suffix == ".pdf":
            import pymupdf
            with pymupdf.open(path) as doc:
                if doc.needs_pass or len(doc) > 100:
                    raise ValueError("PDF 加密或超过 100 页，请先拆分/解密后导入")
                pages = [(i+1, page.get_text()) for i, page in enumerate(doc)]
            empty_pages = [str(p) for p, text in pages if not text.strip()]
            if empty_pages:
                raise ValueError("PDF 第 " + ",".join(empty_pages[:10]) + " 页无可提取文字（可能为扫描页）。请本地 OCR 并人工校对后导入 TXT；本版不静默跳过")
            return pages
        raise ValueError("仅支持 PDF / DOCX / UTF-8 TXT / MD")
    except ValueError:
        raise
    except Exception:
        raise ValueError("文档无法解析。请检查格式、编码或文件是否损坏") from None


def chunks(pages):
    for page, text in pages:
        section = "正文"
        for paragraph in re.split(r"\n\s*\n", text):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            if paragraph.startswith("#"):
                section = paragraph.splitlines()[0].lstrip("# ")[:120]
            for offset in range(0, len(paragraph), 650):
                yield page, section, paragraph[offset:offset+650]


def import_document(db, data_dir, path, product_id, kind, provenance, source_url="", version="1",
                    effective_at="", expires_at=""):
    path = Path(path)
    product = Catalog(db).product(product_id)
    if not product or kind not in KINDS or not provenance.strip():
        raise ValueError("需选择商品、资料类别并填写来源说明")
    if not path.is_file() or path.stat().st_size > 20*1024*1024:
        raise ValueError("文件不存在或超过 20 MB")
    if kind == "shop_policy":
        try:
            if date.fromisoformat(effective_at) > date.fromisoformat(expires_at):
                raise ValueError()
        except ValueError:
            raise ValueError("SOP 必须填写正确的生效/失效日期 YYYY-MM-DD") from None
    # Parse the exact bytes copied, not a file that may change between hash and parse.
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    copy_dir = Path(data_dir) / "knowledge"
    copy_dir.mkdir(parents=True, exist_ok=True)
    copy = copy_dir / (digest + path.suffix.lower())
    if not copy.exists():
        copy.write_bytes(raw)
    parts = list(chunks(extract_pages(copy)))
    if not parts:
        raise ValueError("文档没有可用正文，请人工检查")
    if sum(len(t) for _, _, t in parts) > 1_000_000:
        raise ValueError("正文超过一百万字，请拆分后导入")
    duplicate = db.rows("SELECT id FROM knowledge_sources WHERE product_id=? AND sha256=? AND kind=? AND version=?",
                        (product_id, digest, kind, version))
    if duplicate:
        raise ValueError("同型号的相同文件版本已存在，未重复导入")
    source_id = uuid4().hex
    with db.connect() as conn:
        conn.execute("""INSERT INTO knowledge_sources
          (id,product_id,kind,title,path,source_url,provenance,sha256,version,is_demo,imported_at,effective_at,expires_at)
          VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""", (source_id, product_id, kind, path.name, str(copy.resolve()),
          source_url, provenance, digest, version, product["is_demo"], now(), effective_at, expires_at))
        conn.executemany("INSERT INTO knowledge_chunks VALUES (?,?,?,?,?)",
                         [(uuid4().hex, source_id, section, page, text) for page, section, text in parts])
    return source_id
